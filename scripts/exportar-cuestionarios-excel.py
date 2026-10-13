# -*- coding: utf-8 -*-
"""Arma el Excel con los cuestionarios tal como están publicados, para el cliente."""
import io
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

ORIGEN, DESTINO = sys.argv[1], sys.argv[2]

# ── Paleta: dorado de la marca sobre gris claro ──
DORADO = 'C8A04A'
GRIS = 'EFEFEF'
GRIS_SUAVE = 'F8F8F8'

T_HOJA = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
T_SECCION = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
T_SUB = Font(name='Calibri', size=11, bold=True)
T_NORMAL = Font(name='Calibri', size=11)
T_TENUE = Font(name='Calibri', size=10, italic=True, color='666666')

F_HOJA = PatternFill('solid', fgColor='1A1A1A')
F_SECCION = PatternFill('solid', fgColor=DORADO)
F_SUB = PatternFill('solid', fgColor=GRIS)
F_ALT = PatternFill('solid', fgColor=GRIS_SUAVE)

ARRIBA = Alignment(vertical='top', wrap_text=True)
ARRIBA_CENTRO = Alignment(vertical='top', horizontal='center', wrap_text=True)
CENTRO = Alignment(vertical='center', wrap_text=True)

_linea = Side(style='thin', color='DDDDDD')
BORDE = Border(left=_linea, right=_linea, top=_linea, bottom=_linea)

ANCHOS = {'A': 15, 'B': 40, 'C': 48, 'D': 48}
COLUMNAS = len(ANCHOS)
LINEA = 14.5  # alto de un renglón en Calibri 11

NOMBRE_HOJA = {
    'emprendimiento': 'Emprendimiento',
    'lider': 'Líder',
    'evaluador': 'Líder Evaluador',
}
NIVEL = {'alto': 'Alto', 'medio': 'Medio', 'bajo': 'Bajo', 'muyBajo': 'Muy bajo'}
NORTE = {'aceleracion': 'Norte A', 'estructura': 'Norte B'}


class Hoja:
    """
    Escribe de arriba hacia abajo llevando la cuenta de la fila.

    Cada fila se reparte en cuatro columnas según los `tramos` que se le pasen:
    `tramos=(1, 1, 2)` deja la tercera celda ocupando C y D. Sin esto las tablas
    de dos o tres columnas dejaban una columna vacía al borde de la hoja.
    """

    def __init__(self, ws):
        self.ws = ws
        self.f = 1

    def _pintar(self, col, tramo, valor, fuente, relleno, alineacion):
        if tramo > 1:
            self.ws.merge_cells(start_row=self.f, start_column=col,
                                end_row=self.f, end_column=col + tramo - 1)
        # El estilo va en todas las celdas del rango: en una combinada, el borde
        # derecho lo dibuja la última celda, no la que tiene el texto.
        for i in range(tramo):
            c = self.ws.cell(self.f, col + i)
            c.font = fuente
            c.alignment = alineacion
            c.border = BORDE
            if relleno:
                c.fill = relleno
        self.ws.cell(self.f, col, valor)

    def _escribir(self, valores, tramos, fuente, relleno, centrar_primera=False):
        col = 1
        for i, (v, t) in enumerate(zip(valores, tramos)):
            # Sólo se centra la primera celda cuando es la columna angosta
            # (número de afirmación, rango de puntaje), no un texto en A:B.
            corta = i == 0 and t == 1 and len(str(v)) <= 12
            alineacion = ARRIBA_CENTRO if (centrar_primera or corta) else ARRIBA
            self._pintar(col, t, v, fuente, relleno, alineacion)
            col += t
        self.f += 1

    def titulo(self, texto, subtitulo=''):
        self._pintar(1, COLUMNAS, texto, T_HOJA, F_HOJA, CENTRO)
        self.ws.row_dimensions[self.f].height = 30
        self.f += 1
        if subtitulo:
            self._pintar(1, COLUMNAS, subtitulo, T_TENUE, None, CENTRO)
            self.f += 1
        self.f += 1

    def seccion(self, texto):
        self._pintar(1, COLUMNAS, texto, T_SECCION, F_SECCION, CENTRO)
        self.ws.row_dimensions[self.f].height = 22
        self.f += 1

    def subseccion(self, texto):
        self._pintar(1, COLUMNAS, texto, T_SUB, F_SUB, CENTRO)
        self.f += 1

    def encabezados(self, titulos, tramos):
        self._escribir(titulos, tramos, T_SUB, F_SUB, centrar_primera=True)

    def fila(self, valores, tramos, alterna=False):
        self._escribir(valores, tramos, T_NORMAL, F_ALT if alterna else None)

    def tabla(self, titulos, tramos, filas):
        self.encabezados(titulos, tramos)
        for i, vals in enumerate(filas):
            self.fila(vals, tramos, alterna=i % 2 == 1)

    def parrafo(self, texto, fuente=None):
        self._pintar(1, COLUMNAS, texto, fuente or T_NORMAL, None, ARRIBA)
        self.f += 1

    def blanco(self, n=1):
        self.f += n


def ajustar_alturas(ws):
    """
    Calcula la altura de cada fila según el texto que tiene.

    Excel auto-ajusta las filas con texto que envuelve, pero **no lo hace cuando
    la celda está combinada**: la intro y la lectura general se verían cortadas
    en un solo renglón. Como hay que calcularlas igual, se calculan todas y así
    el archivo se ve igual en Excel, LibreOffice y la vista previa del correo.
    """
    # Ancho real de cada celda combinada, por su esquina superior izquierda.
    ancho_combinado = {}
    for r in ws.merged_cells.ranges:
        ancho = sum(ANCHOS[get_column_letter(c)]
                    for c in range(r.min_col, r.max_col + 1))
        ancho_combinado[(r.min_row, r.min_col)] = ancho

    for f in range(1, ws.max_row + 1):
        # Las filas con altura fija (título, secciones) ya están resueltas.
        if f in ws.row_dimensions and ws.row_dimensions[f].height:
            continue
        renglones = 1
        for col in range(1, COLUMNAS + 1):
            valor = ws.cell(f, col).value
            if not isinstance(valor, str) or not valor:
                continue
            ancho = ancho_combinado.get((f, col), ANCHOS[get_column_letter(col)])
            # En Calibri 11 entran ~1,15 caracteres por unidad de ancho; se usa
            # 1,1 para dejar margen y que ninguna fila quede corta.
            por_linea = max(10, int(ancho * 1.1) - 2)
            alto = sum(max(1, -(-len(p) // por_linea)) for p in valor.split('\n'))
            renglones = max(renglones, alto)
        ws.row_dimensions[f].height = round(renglones * LINEA + 4, 1)


def texto_banda(b):
    """Las bandas traen un párrafo o una lista de viñetas, según el cuestionario."""
    if b.get('bullets'):
        return '\n'.join('• ' + x for x in b['bullets'])
    return b.get('text', '')


def volcar(ws, a):
    h = Hoja(ws)
    preguntas = sum(len(p['questions']) for p in a['pilares'])
    h.titulo(a['title'], '%d afirmaciones · %d pilares · escala 1 a 5 · puntaje de %d a %d'
             % (preguntas, len(a['pilares']), preguntas, preguntas * 5))

    # ── Presentación y datos previos ──
    h.seccion('Texto de presentación')
    h.parrafo(a['intro'])
    h.blanco()

    h.seccion(a['orgFieldsTitle'])
    h.tabla(['Dato que se pide', 'Ejemplo / opciones'], (2, 2),
            [[f['label'] + (' (opcional)' if f.get('optional')
                            and 'opcional' not in f['label'] else ''),
              ' · '.join(f['options']) if f.get('options') else f.get('placeholder', '')]
             for f in a['orgFields']])
    h.blanco()

    h.seccion('Escala de respuesta')
    h.tabla(['Valor', 'Etiqueta'], (1, 3),
            [[s['value'], s['label']] for s in a['scale']])
    h.blanco()

    # ── Pilares: preguntas y devolución por banda ──
    for n, p in enumerate(a['pilares'], start=1):
        h.seccion('Pilar %d: %s' % (n, p['name']))
        # La columna de puntaje invertido sólo aparece si el pilar la usa.
        if any(q.get('reverse') for q in p['questions']):
            h.tabla(['N°', 'Afirmación', 'Puntaje invertido'], (1, 2, 1),
                    [[q['id'], q['text'], 'Sí' if q.get('reverse') else '']
                     for q in p['questions']])
        else:
            h.tabla(['N°', 'Afirmación'], (1, 3),
                    [[q['id'], q['text']] for q in p['questions']])
        h.blanco()
        h.subseccion('Devolución del Pilar %d según el puntaje' % n)
        h.tabla(['Puntaje', 'Título', 'Devolución'], (1, 1, 2),
                [['%d a %d' % (b['min'], b['max']), b['title'], texto_banda(b)]
                 for b in p['bands']])
        h.blanco()

    # ── Niveles globales ──
    h.seccion('Niveles según el puntaje total')
    h.tabla(['Puntaje', 'Nivel', 'Devolución', 'Cierre'], (1, 1, 1, 1),
            [['%d a %d' % (g['min'], g['max']),
              g['title'] + ('\n\nPerfil: ' + g['perfil'] if g.get('perfil') else ''),
              g.get('text', ''), g.get('closing', '')]
             for g in a['globalBands']])
    h.blanco()

    # ── Norte ──
    h.seccion(a['norteQuestion'])
    h.tabla(['Opción', 'Texto'], (1, 3),
            [[NORTE.get(o['key'], o['key']), o['label']] for o in a['norteOptions']])
    h.blanco()

    # ── Cruce norte × nivel ──
    if a.get('cruces'):
        h.seccion('Cierre según el norte elegido y el nivel alcanzado')
        h.tabla(['Norte y nivel', 'Lectura de cruce', 'Diagnóstico y primer paso'], (1, 1, 2),
                [['%s\nNivel %s' % (NORTE.get(c['norteKey'], c['norteKey']),
                                    NIVEL.get(c['level'], c['level']).lower()),
                  c.get('lectura', ''),
                  'Diagnóstico: %s\n\nPrimer paso sugerido: %s'
                  % (c.get('diagnostico', ''), c.get('primerPaso', ''))]
                 for c in a['cruces']])
        h.blanco()

    # ── Lectura general ──
    if a.get('lecturaGeneral'):
        h.seccion('Lectura general (va en todos los informes)')
        for t in a['lecturaGeneral']:
            h.parrafo(t)
        h.blanco()

    # ── Ancho de columnas, alturas y vista ──
    for col, ancho in ANCHOS.items():
        ws.column_dimensions[col].width = ancho
    ajustar_alturas(ws)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A4'
    # Para imprimir o pasar a PDF: a lo ancho de una hoja A4 apaisada.
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:1'


def main():
    datos = json.load(io.open(ORIGEN, encoding='utf-8'))
    wb = Workbook()
    wb.remove(wb.active)

    for a in datos:
        ws = wb.create_sheet(NOMBRE_HOJA.get(a['key'], a['key'])[:31])
        volcar(ws, a)
        print('hoja "%s": %d filas' % (ws.title, ws.max_row))

    wb.save(DESTINO)
    print('\nguardado:', DESTINO)


main()
