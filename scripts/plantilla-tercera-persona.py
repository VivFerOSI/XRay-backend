# -*- coding: utf-8 -*-
"""
Arma la plantilla que el cliente completa con las afirmaciones en tercera
persona, necesarias para la evaluación 360°.

En el 360° quien responde habla de otra persona, así que cada afirmación
necesita su versión reescrita. Sin ellas `listoPara360()` devuelve false y el
sitio ni ofrece invitar evaluadores.

    node scripts/exportar-definiciones.js definiciones.json
    python scripts/plantilla-tercera-persona.py definiciones.json "Plantilla.xlsx"
"""
import io
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

ORIGEN, DESTINO = sys.argv[1], sys.argv[2]

DORADO = 'C8A04A'
NAVY = '1A1A1A'

T_TITULO = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
T_SECCION = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
T_CABECERA = Font(name='Calibri', size=11, bold=True)
T_NORMAL = Font(name='Calibri', size=11)
T_EJEMPLO = Font(name='Calibri', size=11, italic=True, color='7A5C1E')
T_AYUDA = Font(name='Calibri', size=10, color='555555')

F_TITULO = PatternFill('solid', fgColor=NAVY)
F_SECCION = PatternFill('solid', fgColor=DORADO)
F_CABECERA = PatternFill('solid', fgColor='EFEFEF')
F_EJEMPLO = PatternFill('solid', fgColor='FDF6E3')
F_COMPLETAR = PatternFill('solid', fgColor='FFFDF5')

ARRIBA = Alignment(vertical='top', wrap_text=True)
CENTRO = Alignment(vertical='center', horizontal='center', wrap_text=True)
_l = Side(style='thin', color='DDDDDD')
BORDE = Border(left=_l, right=_l, top=_l, bottom=_l)

# Sólo los cuestionarios que tienen el 360° habilitado.
NOMBRE = {
    'lider': 'Persona Evaluada',
    'evaluador': 'Evaluador',
}

AYUDA = (
    'Completá la última columna con la misma afirmación, pero hablando de la persona '
    'evaluada en lugar de uno mismo. En general alcanza con pasar el verbo a tercera '
    'persona y cambiar "mi/mis" por "su/sus". Si alguna frase necesita más vueltas, '
    'escribila como te suene natural: lo importante es que quien responde esté '
    'describiendo a otra persona.'
)

# El ejemplo va escrito a mano y no generado: es el modelo que se va a copiar
# 50 veces, así que no puede salir de una regla que confunda "comprende" con
# "comprenda".
EJEMPLOS = {
    'lider': (
        'Comprendo con claridad los objetivos estratégicos de la organización y puedo '
        'explicárselos a mi equipo de manera sencilla.',
        'Comprende con claridad los objetivos estratégicos de la organización y puede '
        'explicárselos a su equipo de manera sencilla.',
    ),
    'evaluador': (
        'Registro de forma periódica hechos concretos, logros y evidencias observables '
        'de la persona evaluada durante todo el año.',
        'Registra de forma periódica hechos concretos, logros y evidencias observables '
        'de la persona evaluada durante todo el año.',
    ),
}


def volcar(ws, a):
    fila = 1

    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
    c = ws.cell(fila, 1, 'Evaluación 360° — %s' % NOMBRE[a['key']])
    c.font, c.fill, c.alignment = T_TITULO, F_TITULO, CENTRO
    ws.row_dimensions[fila].height = 28
    fila += 2

    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
    c = ws.cell(fila, 1, AYUDA)
    c.font, c.alignment = T_AYUDA, ARRIBA
    ws.row_dimensions[fila].height = 46
    fila += 2

    # Ejemplo resuelto, para que se vea de un vistazo qué se espera.
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
    c = ws.cell(fila, 1, 'Ejemplo')
    c.font, c.fill, c.alignment = T_SECCION, F_SECCION, CENTRO
    fila += 1
    for col, txt in enumerate(
        ['', 'Así está hoy', 'Así quedaría', ''], start=1
    ):
        c = ws.cell(fila, col, txt)
        c.font, c.fill, c.alignment, c.border = T_CABECERA, F_CABECERA, CENTRO, BORDE
    fila += 1
    actual, en_tercera = EJEMPLOS[a['key']]
    for col, txt in enumerate(['', actual, en_tercera, ''], start=1):
        c = ws.cell(fila, col, txt)
        c.font, c.fill, c.alignment, c.border = T_EJEMPLO, F_EJEMPLO, ARRIBA, BORDE
    ws.row_dimensions[fila].height = 46
    fila += 2

    # Las 25 afirmaciones, agrupadas por pilar.
    for n, p in enumerate(a['pilares'], start=1):
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
        c = ws.cell(fila, 1, 'Pilar %d: %s' % (n, p['name']))
        c.font, c.fill, c.alignment = T_SECCION, F_SECCION, CENTRO
        fila += 1

        for col, txt in enumerate(
            ['N°', 'Afirmación actual', 'Afirmación en tercera persona', 'Notas'],
            start=1,
        ):
            c = ws.cell(fila, col, txt)
            c.font, c.fill, c.alignment, c.border = T_CABECERA, F_CABECERA, CENTRO, BORDE
        fila += 1

        for q in p['questions']:
            ws.cell(fila, 1, q['id']).alignment = CENTRO
            ws.cell(fila, 2, q['text']).alignment = ARRIBA
            # Esta es la que se completa: queda vacía y con fondo propio.
            c = ws.cell(fila, 3, '')
            c.fill, c.alignment = F_COMPLETAR, ARRIBA
            ws.cell(fila, 4, '').alignment = ARRIBA
            for col in range(1, 5):
                celda = ws.cell(fila, col)
                celda.border = BORDE
                if celda.font.sz is None:
                    celda.font = T_NORMAL
            ws.row_dimensions[fila].height = 46
            fila += 1
        fila += 1

    for col, ancho in zip('ABCD', (6, 58, 58, 26)):
        ws.column_dimensions[col].width = ancho
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A3'
    # Para imprimir o pasar a PDF: las dos columnas de texto en una misma hoja.
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    datos = json.load(io.open(ORIGEN, encoding='utf-8'))
    wb = Workbook()
    wb.remove(wb.active)

    for a in datos:
        if a['key'] not in NOMBRE:
            continue
        ws = wb.create_sheet(NOMBRE[a['key']])
        volcar(ws, a)
        preguntas = sum(len(p['questions']) for p in a['pilares'])
        print('hoja "%s": %d afirmaciones para completar' % (ws.title, preguntas))

    wb.save(DESTINO)
    print('\nguardado:', DESTINO)


main()
